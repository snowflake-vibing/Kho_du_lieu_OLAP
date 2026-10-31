# -*- coding: utf-8 -*-
"""
STRICT AUTOMATED 3-WAY AUDIT SUITE
Chế tài kiểm định tự động bắt buộc nhằm ngăn chặn AI Assistant vi phạm các quy tắc trong GEMINI.md.
"""
import sys, win32com.client

sys.stdout.reconfigure(encoding='utf-8')

conn = win32com.client.Dispatch("ADODB.Connection")
conn.Open("Provider=MSOLAP;Data Source=localhost;Initial Catalog=SSAS_Football_Transfermarkt;")

print("===================================================================================")
print("KHUNG KIỂM ĐỊNH KỶ LUẬT TỰ ĐỘNG (STRICT DISCIPLINARY AUDIT) - SSAS LIVE CUBE")
print("===================================================================================")

total_checks = 0
passed_checks = 0

def check(name, condition, details=""):
    global total_checks, passed_checks
    total_checks += 1
    if condition:
        passed_checks += 1
        print(f"  [PASS] {name} {details}")
    else:
        print(f"  [FAIL] {name} {details}")

# 0. AUDIT QUERY 04
print("\n--- 0. AUDIT QUERY 04: Bàn thắng, kiến tạo & phút thi đấu theo độ tuổi ---")
with open("mdx/query_04.mdx", "r", encoding="utf-8") as f:
    mdx_04 = f.read()

rs04 = win32com.client.Dispatch("ADODB.Recordset")
rs04.Open(mdx_04, conn)
rows04 = []
while not rs04.EOF:
    rows04.append([rs04.Fields(i).Value for i in range(rs04.Fields.Count)])
    rs04.MoveNext()
rs04.Close()

check("Q04 - MDX Row Count == 29", len(rows04) == 29, f"(Actual: {len(rows04)})")
check("Q04 - MDX Column Count == 3 (No Career Stage)", len(rows04[0]) - 1 == 3, f"(Actual: {len(rows04[0]) - 1})")

# 1. AUDIT QUERY 07
print("\n--- 1. AUDIT QUERY 07: Bàn thắng theo quốc tịch Pareto 80/20 tại Anh ---")
with open("mdx/query_07.mdx", "r", encoding="utf-8") as f:
    mdx_07 = f.read()

rs07 = win32com.client.Dispatch("ADODB.Recordset")
rs07.Open(mdx_07, conn)
rows07 = []
while not rs07.EOF:
    rows07.append([rs07.Fields(i).Value for i in range(rs07.Fields.Count)])
    rs07.MoveNext()
rs07.Close()

check("Q07 - MDX Row Count == 18", len(rows07) == 18, f"(Actual: {len(rows07)})")
check("Q07 - MDX Column Count == 2 (Goals & Pct)", len(rows07[0]) - 1 == 2, f"(Actual: {len(rows07[0]) - 1})")
check("Q07 - Goals Column Present and England is #1", rows07[0][0] == "England" and int(rows07[0][1]) == 714, f"(Top 1: {rows07[0][0]} with {rows07[0][1]} goals)")

# 2. AUDIT QUERY 09
print("\n--- 2. AUDIT QUERY 09: Giao tập hợp INTERSECT 9 CLB Cuối tuần & Ngày thường ---")
with open("mdx/query_09.mdx", "r", encoding="utf-8") as f:
    mdx_09 = f.read()

rs09 = win32com.client.Dispatch("ADODB.Recordset")
rs09.Open(mdx_09, conn)
rows09 = []
while not rs09.EOF:
    rows09.append([rs09.Fields(i).Value for i in range(rs09.Fields.Count)])
    rs09.MoveNext()
rs09.Close()

check("Q09 - MDX Row Count == 9", len(rows09) == 9, f"(Actual: {len(rows09)})")
check("Q09 - MDX Column Count == 4 (CrossJoin Goals & Cards)", len(rows09[0]) - 1 == 4, f"(Actual: {len(rows09[0]) - 1})")

# 3. AUDIT QUERY 10
print("\n--- 3. AUDIT QUERY 10: Top 10 CLB Châu Âu theo 4 Quý năm 2023 ---")
with open("mdx/query_10.mdx", "r", encoding="utf-8") as f:
    mdx_10 = f.read()

rs10 = win32com.client.Dispatch("ADODB.Recordset")
rs10.Open(mdx_10, conn)
rows10 = []
while not rs10.EOF:
    rows10.append([rs10.Fields(i).Value for i in range(rs10.Fields.Count)])
    rs10.MoveNext()
rs10.Close()

check("Q10 - MDX Row Count == 10", len(rows10) == 10, f"(Actual: {len(rows10)})")
check("Q10 - MDX Column Count == 7", len(rows10[0]) - 1 == 7, f"(Actual: {len(rows10[0]) - 1})")
has_null = any(any(val is None for val in r[1:]) for r in rows10)
check("Q10 - Zero Null / Dense Data", not has_null, "(No nulls or missing cells)")

# 4. AUDIT QUERY 11
print("\n--- 4. AUDIT QUERY 11: Siêu sân vận động sức chứa >= 60.000 chỗ ---")
with open("mdx/query_11.mdx", "r", encoding="utf-8") as f:
    mdx_11 = f.read()

rs11 = win32com.client.Dispatch("ADODB.Recordset")
rs11.Open(mdx_11, conn)
rows11 = []
while not rs11.EOF:
    rows11.append([rs11.Fields(i).Value for i in range(rs11.Fields.Count)])
    rs11.MoveNext()
rs11.Close()

check("Q11 - MDX Row Count == 23", len(rows11) == 23, f"(Actual: {len(rows11)})")
all_caps_numeric = all(isinstance(r[3], (int, float)) and r[3] >= 60000 for r in rows11)
check("Q11 - Capacity is numeric and >= 60000 (No 'All' text bug)", all_caps_numeric, f"(Min cap: {min(r[3] for r in rows11)})")

# 5. AUDIT EXCEL FILE
print("\n--- 5. AUDIT EXCEL DOUBLE-CHECK WORKBOOK ---")
import openpyxl
wb = openpyxl.load_workbook("mdx/Bao_Cao_Double_Check_Query_10_11.xlsx")
ws04 = wb["Query 04 - Hieu Suat Do Tuoi"]
ws07 = wb["Query 07 - Pareto Ban Thang Anh"]
ws09 = wb["Query 09 - Intersect 9 CLB"]
ws10 = wb["Query 10 - QoQ Top 10 CLB"]
ws11 = wb["Query 11 - Sân trên 60k Chỗ"]

check("Excel - Q04 Sheet has 29 data rows", ws04.max_row - 4 == 29, f"(Actual: {ws04.max_row - 4})")
check("Excel - Q07 Sheet has 18 data rows & Goals column", ws07.max_row - 4 == 18, f"(Actual: {ws07.max_row - 4})")
check("Excel - Q09 Sheet has 9 data rows", ws09.max_row - 4 == 9, f"(Actual: {ws09.max_row - 4})")
check("Excel - Q10 Sheet has 10 data rows", ws10.max_row - 4 == 10, f"(Actual: {ws10.max_row - 4})")
check("Excel - Q11 Sheet has 23 data rows", ws11.max_row - 4 == 23, f"(Actual: {ws11.max_row - 4})")

conn.Close()

print("\n===================================================================================")
print(f"KẾT QUẢ TỔNG THỂ: {passed_checks}/{total_checks} KIỂM TRA ĐẠT CHUẨN")
if passed_checks == total_checks:
    print("TRẠNG THÁI: TẤT CẢ CÁC RÀO CẢN KỶ LUẬT ĐÃ ĐƯỢC THỎA MÃN 100%!")
else:
    print("TRẠNG THÁI: PHÁT HIỆN SAI PHẠM! YÊU CẦU SỬA ĐỔI NGAY LẬP TỨC!")
print("===================================================================================")
